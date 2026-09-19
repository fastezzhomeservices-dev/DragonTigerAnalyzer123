import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from datetime import datetime
from collections import Counter
import csv
import re
import html
import os

# Windows HD / High-DPI rendering: keep Tkinter text and controls sharp instead of bitmap-scaled.
if os.name == 'nt':
    try:
        import ctypes
        _user32 = ctypes.windll.user32
        _user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4))  # Per-monitor DPI aware V2
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass
import json
import shutil
import webbrowser
import threading
from PIL import Image, ImageTk

CARDS = ['A','2','3','4','5','6','7','8','9','10','J','Q','K']
ODD = {'A','3','5','7','9','J','K'}
RESULTS = ['D','T','TIE']
RESULT_COLORS = {'D':'#00a8ff','T':'#ff2028','TIE':'#00e676'}


def clean_card(v):
    v = str(v or '').strip().upper()
    v = re.sub(r'[^A-Z0-9]', '', v)
    if v in CARDS:
        return v
    return ''


def normalize_result(v):
    v = str(v or '').strip().upper()
    if v in ('DRAGON','D'):
        return 'D'
    if v in ('TIGER','T'):
        return 'T'
    if v in ('TIE','DRAW','SAME'):
        return 'TIE'
    return ''


class App:
    def __init__(self, root):
        self.root = root
        root.title('Dragon Tiger Analyzer - PC v2.0')
        root.geometry('1280x720')
        root.minsize(1100, 650)
        root.configure(bg='#111827')
        self.data = []
        self.pair_history = []
        self.result_history = []
        self.current_prediction = tk.StringVar(value='')
        self.prediction_pct = tk.StringVar(value='')
        self.prediction_correct = tk.StringVar(value='')
        self.actual_dragon = tk.StringVar(value='')
        self.actual_tiger = tk.StringVar(value='')
        self.final_result = tk.StringVar(value='D')
        self.dr = tk.StringVar(value='A')
        self.ti = tk.StringVar(value='A')
        self.re = tk.StringVar(value='D')
        self.pair = tk.StringVar(value='JQ')
        self.summary = tk.StringVar(value='PAIR JQ | NO HISTORY')
        self.collect_date = tk.StringVar(value=datetime.now().strftime('%d/%m/%Y'))
        self.entries = tk.IntVar(value=50)
        self.status = tk.StringVar(value='Ready')
        self.settings_dir = os.path.join(os.environ.get('APPDATA', os.path.expanduser('~')), 'DragonTigerAnalyzer')
        self.settings_file = os.path.join(self.settings_dir, 'settings.json')
        self.theme_name = 'Midnight Blue'
        self.background_path = ''
        self.bg_photo = None
        self.live_url = 'https://games.evolution.com/live-casino/dragon-tiger/'
        self.live_open = False
        self.live_window = None
        self.load_settings()
        self.build_styles()
        # Create the background canvas before all UI widgets so no widget-level lower() call is needed.
        self.setup_background()
        self.build_ui()
        self.refresh()

    def load_settings(self):
        try:
            with open(self.settings_file, 'r', encoding='utf-8') as f:
                s = json.load(f)
            self.theme_name = s.get('theme', 'Midnight Blue')
            if self.theme_name == 'Dark Red':
                self.theme_name = 'Midnight Blue'
            self.background_path = s.get('background', '')
            self.pair_history = s.get('pair_history', []) if isinstance(s.get('pair_history', []), list) else []
            self.result_history = s.get('result_history', []) if isinstance(s.get('result_history', []), list) else []
            self.pair_history = self.pair_history[-100:]
            self.result_history = self.result_history[-200:]
        except Exception:
            pass

    def save_settings(self):
        try:
            os.makedirs(self.settings_dir, exist_ok=True)
            with open(self.settings_file, 'w', encoding='utf-8') as f:
                json.dump({'theme': self.theme_name, 'background': self.background_path,
                       'pair_history': self.pair_history[-100:],
                       'result_history': self.result_history[-200:]}, f, indent=2)
        except Exception:
            pass

    def setup_background(self):
        self.bg_canvas = tk.Canvas(self.root, highlightthickness=0, bd=0)
        self.bg_canvas.place(x=0, y=0, relwidth=1, relheight=1)
        # Background canvas is created before the other UI widgets, so it naturally stays behind them.
        self.root.bind('<Configure>', self._resize_background)
        self.apply_theme(self.theme_name, save=False)

    def _resize_background(self, event=None):
        if not getattr(self, 'background_path', '') or not os.path.exists(self.background_path):
            return
        try:
            im = Image.open(self.background_path).convert('RGB')
            w, h = max(1, self.root.winfo_width()), max(1, self.root.winfo_height())
            im.thumbnail((w, h), Image.LANCZOS)
            self.bg_photo = ImageTk.PhotoImage(im)
            self.bg_canvas.delete('all')
            self.bg_canvas.create_image(w//2, h//2, image=self.bg_photo, anchor='center')
        except Exception:
            pass

    def choose_background(self):
        path = filedialog.askopenfilename(filetypes=[('Images','*.png;*.jpg;*.jpeg;*.gif;*.bmp'),('All files','*.*')])
        if not path:
            return
        try:
            os.makedirs(self.settings_dir, exist_ok=True)
            dest = os.path.join(self.settings_dir, 'background' + os.path.splitext(path)[1].lower())
            shutil.copy2(path, dest)
            self.background_path = dest
            self.save_settings()
            self._resize_background()
            self.status.set('Custom background applied and saved.')
        except Exception as e:
            messagebox.showerror('Background', str(e))

    def reset_background(self):
        self.background_path = ''
        self.save_settings()
        self.bg_photo = None
        self.bg_canvas.delete('all')
        self.bg_canvas.configure(bg=self._theme()['root'])
        self.status.set('Background reset to theme default.')

    def export_theme(self):
        path = filedialog.asksaveasfilename(defaultextension='.json', filetypes=[('Theme Settings','*.json')], initialfile='DragonTigerTheme.json')
        if not path:
            return
        try:
            payload = {'theme': self.theme_name}
            if self.background_path and os.path.exists(self.background_path):
                ext = os.path.splitext(self.background_path)[1].lower() or '.png'
                bg_dest = os.path.splitext(path)[0] + ext
                shutil.copy2(self.background_path, bg_dest)
                payload['background_file'] = os.path.basename(bg_dest)
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(payload, f, indent=2)
            messagebox.showinfo('Theme Export', 'Theme settings exported successfully. If a custom background is used, its image was exported beside the JSON file.')
        except Exception as e:
            messagebox.showerror('Theme Export', str(e))

    def _theme(self):
        themes = {
            'Dark Red': {'root':'#111827','panel':'#111827','header':'#7f1d1d','heading':'#7f1d1d','button':'#374151','text':'#f9fafb','accent':'#fbbf24'},
            'Midnight Blue': {'root':'#0b1220','panel':'#0b1220','header':'#1e3a8a','heading':'#1e40af','button':'#1e293b','text':'#f8fafc','accent':'#60a5fa'},
            'Emerald': {'root':'#071a14','panel':'#071a14','header':'#065f46','heading':'#047857','button':'#1f2937','text':'#ecfdf5','accent':'#34d399'},
            'Purple': {'root':'#140d20','panel':'#140d20','header':'#581c87','heading':'#6b21a8','button':'#312e3f','text':'#faf5ff','accent':'#d8b4fe'},
            'Light': {'root':'#eef2f7','panel':'#eef2f7','header':'#334155','heading':'#475569','button':'#dbe2ea','text':'#111827','accent':'#b45309'}
        }
        return themes.get(self.theme_name, themes['Dark Red'])

    def apply_theme(self, name=None, save=True):
        if name:
            self.theme_name = name
        t = self._theme()
        def walk(w):
            try:
                cls = w.winfo_class()
                if cls in ('Frame','Labelframe','TFrame','TLabelframe'):
                    w.configure(bg=t['panel'])
                elif cls in ('Label','TLabel'):
                    w.configure(bg=t['panel'], fg=t['text'])
                elif cls in ('Button','TButton'):
                    w.configure(bg=t['button'], fg=t['text'])
                elif cls == 'Canvas' and w is not getattr(self, 'bg_canvas', None):
                    w.configure(bg=t['panel'])
            except Exception:
                pass
            for child in w.winfo_children():
                walk(child)
        walk(self.root)
        try:
            self.root.configure(bg=t['root'])
            self.bg_canvas.configure(bg=t['root'])
            self.root.children.get('!frame')
            s=ttk.Style()
            s.configure('TFrame', background=t['panel'])
            s.configure('TLabelframe', background=t['panel'], foreground=t['text'])
            s.configure('TLabelframe.Label', background=t['panel'], foreground=t['accent'])
            s.configure('TLabel', background=t['panel'], foreground=t['text'])
            s.configure('TButton', background=t['button'], foreground=t['text'])
            s.map('TButton', background=[('active',t['heading'])])
            s.configure('Treeview', background='#1f2937' if self.theme_name != 'Light' else '#ffffff', fieldbackground='#1f2937' if self.theme_name != 'Light' else '#ffffff', foreground='#f9fafb' if self.theme_name != 'Light' else '#111827')
            s.configure('Treeview.Heading', background=t['heading'], foreground='white')
            s.configure('TCombobox', fieldbackground='#1f2937' if self.theme_name != 'Light' else '#ffffff', background=t['button'], foreground=t['text'], selectbackground=t['heading'], selectforeground='white')
        except Exception:
            pass
        if save:
            self.save_settings()
        if getattr(self, 'background_path', ''):
            self._resize_background()

    def open_theme_settings(self):
        win=tk.Toplevel(self.root)
        win.title('Theme & Background')
        win.geometry('430x260')
        win.resizable(False, False)
        t=self._theme()
        win.configure(bg=t['root'])
        tk.Label(win,text='THEME & BACKGROUND',bg=t['root'],fg=t['accent'],font=('Segoe UI',14,'bold')).pack(pady=(18,12))
        row=tk.Frame(win,bg=t['root']); row.pack(fill='x',padx=25,pady=6)
        tk.Label(row,text='Theme',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left')
        theme_var=tk.StringVar(value=self.theme_name)
        cb=ttk.Combobox(row,textvariable=theme_var,values=['Dark Red','Midnight Blue','Emerald','Purple','Light'],state='readonly',width=20)
        cb.pack(side='right')
        def apply_and_close():
            self.apply_theme(theme_var.get())
            win.destroy()
        tk.Button(win,text='APPLY THEME',command=apply_and_close,bg=t['heading'],fg='white',font=('Segoe UI',10,'bold'),relief='flat',padx=16,pady=7).pack(pady=8)
        tk.Button(win,text='CHOOSE BACKGROUND IMAGE',command=self.choose_background,bg=t['button'],fg=t['text'],font=('Segoe UI',10,'bold'),relief='flat',padx=12,pady=7).pack(pady=4)
        tk.Button(win,text='RESET BACKGROUND',command=self.reset_background,bg=t['button'],fg=t['text'],font=('Segoe UI',10,'bold'),relief='flat',padx=12,pady=7).pack(pady=4)
        tk.Button(win,text='EXPORT THEME',command=self.export_theme,bg=t['button'],fg=t['text'],font=('Segoe UI',10,'bold'),relief='flat',padx=12,pady=7).pack(pady=4)

    def build_styles(self):
        s = ttk.Style()
        s.theme_use('clam')
        s.configure('TFrame', background='#111827')
        s.configure('TLabelframe', background='#111827', foreground='#e5e7eb')
        s.configure('TLabelframe.Label', background='#111827', foreground='#fbbf24', font=('Segoe UI',11,'bold'))
        s.configure('TLabel', background='#111827', foreground='#f3f4f6', font=('Segoe UI',10))
        s.configure('TButton', font=('Segoe UI',10,'bold'), padding=(11,7), background='#374151', foreground='white')
        s.map('TButton', background=[('active','#4b5563')])
        s.configure('TCombobox', fieldbackground='#1f2937', background='#374151', foreground='#f9fafb', selectbackground='#2563eb', selectforeground='white')
        s.map('TCombobox', fieldbackground=[('readonly','#1f2937')], foreground=[('readonly','#f9fafb')], selectbackground=[('readonly','#2563eb')], selectforeground=[('readonly','white')])
        s.configure('Treeview', background='#1f2937', fieldbackground='#1f2937', foreground='#f9fafb', rowheight=32, font=('Segoe UI',9))
        s.configure('Treeview.Heading', background='#7f1d1d', foreground='white', font=('Segoe UI',9,'bold'), padding=7)

    def open_contact_us(self):
        try:
            import webbrowser
            webbrowser.open('https://wa.me/')
        except Exception as e:
            messagebox.showerror('CONTACT US', f'Unable to open WhatsApp: {e}')

    def build_ui(self):
        # FINAL PC LAYOUT — matches the approved FASTEZZ SERVICES reference.
        BG = '#061326'
        PANEL = '#071a35'
        BLUE = '#0758d9'
        BORDER = '#00a8ff'
        YELLOW = '#ffe600'
        WHITE = '#f8fafc'
        GREEN = '#00e676'
        RED = '#ff2028'

        self.root.configure(bg=BG)
        self.root.geometry('1180x680')
        self.root.minsize(1050,640)

        # Header
        header = tk.Frame(self.root,bg='#06183a',height=88,highlightthickness=1,highlightbackground='#0b72ff')
        header.pack(fill='x')
        header.pack_propagate(False)

        # FASTEZZ SERVICES logo mark — light/white background, matching the supplied reference.
        logo = tk.Canvas(header,width=82,height=76,bg='white',highlightthickness=0)
        logo.pack(side='left',padx=(6,1),pady=3)
        logo.create_oval(27,2,55,30,fill='#f95678',outline='#111827',width=1)
        logo.create_oval(7,29,35,57,fill='#10a9ad',outline='#111827',width=1)
        logo.create_oval(47,29,75,57,fill='#ffc33d',outline='#111827',width=1)
        logo.create_oval(27,47,55,75,fill='#10a9ad',outline='#111827',width=1)
        logo.create_arc(7,29,75,75,start=150,extent=190,style='arc',outline='#111827',width=1)
        logo.create_oval(24,25,58,59,fill='white',outline='#111827',width=2)

        # Minimal top header: FASTEZZ only.
        brand=tk.Frame(header,bg='#06183a')
        brand.pack(side='left',fill='both',expand=True)
        tk.Label(brand,text='FASTEZZ',bg='#06183a',fg='white',
                 font=('Segoe UI',30,'bold')).pack(anchor='w',padx=14,pady=(18,0))

        # Navigation
        nav=tk.Frame(self.root,bg=BG); nav.pack(fill='x',padx=18,pady=(4,3))
        nav_buttons=[
            ('Import Excel/CSV',self.import_data),('Export Excel/CSV',self.export_data),
            ('LIVE DRAGON TIGER',self.open_live_browser),('THEME / BACKGROUND',self.open_theme_settings),
            ('PAIR SEARCH',self.open_pair_search),
            ('PAIR HISTORY',self.show_pair_history),
            ('RESULT HISTORY',self.show_result_history),
            ('PRODUCTION RESULT ENTRY',self.open_production_entry),
            ('DOWNLOAD HISTORY',self.download_history)
        ]
        for text,cmd in nav_buttons:
            tk.Button(nav,text=text,command=cmd,bg='#0758d9',fg='white',activebackground='#0b78ff',
                      activeforeground='white',font=('Segoe UI',9,'bold'),relief='groove',bd=1,
                      padx=8,pady=5).pack(side='left',padx=4)

        # Fixed dashboard area — no page scrollbar. All main sections stay fixed on screen.
        content=tk.Frame(self.root,bg=BG)
        content.pack(fill='both',expand=True,padx=0,pady=0)

        # Mandatory Pair Reference / Statistical Prediction row
        refbar=tk.Frame(content,bg=PANEL,highlightthickness=1,highlightbackground=BORDER,height=58)
        refbar.pack(fill='x',padx=12,pady=(1,2))
        refbar.pack_propagate(False)
        tk.Label(refbar,text='PAIR REFERENCE / STATISTICAL PREDICTION',bg=PANEL,fg=YELLOW,
                 font=('Segoe UI',10,'bold')).pack(side='left',padx=(10,14))
        tk.Label(refbar,text='Pair',bg=PANEL,fg=WHITE,font=('Segoe UI',9,'bold')).pack(side='left',padx=(0,6))
        pair_values_ref=[a+b for a in CARDS for b in CARDS]
        self.reference_pair_var=tk.StringVar(value=self.pair.get() or 'JQ')
        ref_combo=ttk.Combobox(refbar,textvariable=self.reference_pair_var,values=pair_values_ref,
                               state='readonly',width=9,font=('Segoe UI',9,'bold'))
        ref_combo.pack(side='left',ipady=3)
        def run_reference_search(event=None):
            p=self.reference_pair_var.get().strip().upper().replace(' ','')
            if p in pair_values_ref:
                self.pair.set(p)
                self.analyze()
                self.main_pair_var.set(p) if hasattr(self,'main_pair_var') else None
                self.refresh_prediction_dashboard()
        tk.Button(refbar,text='ANALYZE',command=run_reference_search,bg='#e55b19',fg='white',
                  activebackground='#f06b22',activeforeground='white',font=('Segoe UI',9,'bold'),
                  relief='groove',bd=1,padx=18,pady=5).pack(side='left',padx=(8,14))
        tk.Label(refbar,textvariable=self.summary,bg=PANEL,fg=YELLOW,font=('Segoe UI',9,'bold'),
                 anchor='w').pack(side='left',fill='x',expand=True)
        tk.Label(refbar,text='Prediction:',bg=PANEL,fg=WHITE,font=('Segoe UI',9,'bold')).pack(side='left',padx=(8,5))
        tk.Label(refbar,textvariable=self.current_prediction,bg=PANEL,fg=YELLOW,font=('Segoe UI',11,'bold'),
                 width=5).pack(side='left',padx=(0,2))
        tk.Label(refbar,textvariable=self.prediction_pct,bg=PANEL,fg=YELLOW,font=('Segoe UI',9,'bold'),
                 width=8).pack(side='left',padx=(0,8))
        ref_combo.bind('<<ComboboxSelected>>',run_reference_search)

        # Pair search history — pair selector is kept only in the mandatory reference row above.
        # This section is reserved for the latest 15 history results, avoiding a duplicate Enter Pair control.
        ph=tk.LabelFrame(content,text='  PAIR SEARCH HISTORY (LATEST 15)  ',bg=PANEL,fg=YELLOW,
                         font=('Segoe UI',12,'bold'),bd=1,relief='groove')
        ph.pack(fill='x',padx=12,pady=2)

        search_row=tk.Frame(ph,bg=PANEL,height=58)
        search_row.pack(fill='x',padx=8,pady=(3,0))
        search_row.pack_propagate(False)

        self.pair_history_frame=tk.Frame(search_row,bg=PANEL,height=54)
        self.pair_history_frame.pack(fill='both',expand=True)
        self.pair_history_frame.pack_propagate(False)
        self.render_pair_history()

        # Frequency chart
        chart_box=tk.LabelFrame(content,text='  NUMBER / DRAGON-TIGER FREQUENCY CHART  ',bg=PANEL,fg=YELLOW,
                                font=('Segoe UI',12,'bold'),bd=1,relief='groove')
        chart_box.pack(fill='x',padx=12,pady=2)
        self.chart_frame=tk.Frame(chart_box,bg=PANEL,height=28)
        self.chart_frame.pack(fill='x',padx=10,pady=3)
        self.chart_frame.pack_propagate(False)

        # Occurrence report — same table layout, with compact coloured result circles.
        report=tk.LabelFrame(content,text='  OCCURRENCE + PREVIOUS/NEXT 3-ROUND REPORT  ',bg=PANEL,fg=YELLOW,
                             font=('Segoe UI',12,'bold'),bd=1,relief='groove')
        report.pack(fill='x',padx=12,pady=2)
        report_table=tk.Frame(report,bg=PANEL,height=112)
        report_table.pack(fill='x',padx=8,pady=3)
        report_table.pack_propagate(False)

        rcols=('ROUND ID','PREVIOUS 3','MATCH','NEXT 1','NEXT 2','NEXT 3','NEXT 4','NEXT 5','NEXT 6')
        rwidths=[90,250,90,90,90,90,90,90,90]
        self.report_rows=[]

        report_header=tk.Frame(report_table,bg='#0758d9',height=28)
        report_header.pack(fill='x')
        report_header.pack_propagate(False)
        for i,(col,w) in enumerate(zip(rcols,rwidths)):
            report_header.grid_columnconfigure(i,weight=w)
            tk.Label(report_header,text=col,bg='#0758d9',fg='white',
                     font=('Segoe UI',8,'bold'),anchor='center').grid(row=0,column=i,sticky='nsew',padx=1,pady=1)

        report_body=tk.Frame(report_table,bg=PANEL)
        report_body.pack(fill='both',expand=True)
        self.report_canvas=tk.Canvas(report_body,bg=PANEL,highlightthickness=0)
        report_scroll=ttk.Scrollbar(report_body,orient='vertical',command=self.report_canvas.yview)
        self.report_canvas.configure(yscrollcommand=report_scroll.set)
        report_scroll.pack(side='right',fill='y')
        self.report_canvas.pack(side='left',fill='both',expand=True)
        self.report_canvas.bind('<Configure>',lambda e:self.render_report_rows())
        self.render_report_rows()

        # Production result entry table
        prod=tk.LabelFrame(content,text='  PRODUCTION RESULT ENTRY  ',bg=PANEL,fg=YELLOW,
                           font=('Segoe UI',12,'bold'),bd=1,relief='groove')
        prod.pack(fill='x',padx=12,pady=2)
        prod_table=tk.Frame(prod,bg=PANEL,height=68)
        prod_table.pack(fill='x',padx=8,pady=3)
        prod_table.pack_propagate(False)
        pcols=('S NO','ROUND ID','TIME','DRAGON','TIGER','TIE','PREDICTION','RESULT','PAIR','D Q/E','T Q/E','PREV RESULT')
        self.production_tree=ttk.Treeview(prod_table,columns=pcols,show='headings',height=2)
        pwidths={'S NO':60,'ROUND ID':145,'TIME':120,'DRAGON':75,'TIGER':75,'TIE':60,'PREDICTION':105,
                 'RESULT':90,'PAIR':80,'D Q/E':90,'T Q/E':90,'PREV RESULT':110}
        for c in pcols:
            self.production_tree.heading(c,text=c)
            self.production_tree.column(c,width=pwidths[c],anchor='center',stretch=True)
        self.production_tree.pack(fill='both',expand=True)
        self.production_tree.tag_configure('match',background='#078b3d',foreground='white')
        self.production_tree.tag_configure('nomatch',background='#e00012',foreground='white')

        # Bottom prediction history summary
        dash=tk.LabelFrame(content,text='  PREDICTION HISTORY  ',bg=PANEL,fg=YELLOW,
                           font=('Segoe UI',12,'bold'),bd=1,relief='groove')
        dash.pack(fill='x',padx=12,pady=(2,4),ipady=1)
        self.prediction_dashboard=tk.Frame(dash,bg=PANEL)
        self.prediction_dashboard.pack(fill='x',padx=8,pady=2)

        # Hidden compatibility widgets used by existing analysis/export methods.
        # IMPORTANT: keep the single Pair/Summary StringVars created in __init__
        # so the visible Pair Summary and reference selector stay synchronized.
        hidden=tk.Frame(self.root,bg=BG)
        self.hidden_pair=hidden
        pair_combo=ttk.Combobox(hidden,textvariable=self.pair,values=[a+b for a in CARDS for b in CARDS],width=12)
        pair_combo.pack()
        tk.Button(hidden,text='ANALYZE',command=self.analyze).pack()
        self.prediction_label=tk.Label(hidden,textvariable=self.current_prediction)
        self.tree=ttk.Treeview(hidden,columns=('S NO','ROUND ID','TIME','DRAGON','TIGER','RESULT','PAIR','D O/E','T O/E','PREV RESULT'),show='headings')
        self.last_frame=tk.Frame(hidden)
        self.last_analysis=tk.StringVar(value='')
        self.refresh_prediction_dashboard()
        self._render_production_table()

    def open_pair_search(self):
        win=tk.Toplevel(self.root)
        win.title('Pair Search')
        win.geometry('430x220')
        win.resizable(False,False)
        t=self._theme()
        win.configure(bg=t['root'])
        tk.Label(win,text='PAIR SEARCH',bg=t['root'],fg=t['accent'],font=('Segoe UI',14,'bold')).pack(pady=(15,12))
        row=tk.Frame(win,bg=t['root']); row.pack(pady=5)
        tk.Label(row,text='Dragon',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left',padx=(0,5))
        dvar=tk.StringVar(value='A')
        ttk.Combobox(row,textvariable=dvar,values=CARDS,state='readonly',width=5).pack(side='left')
        tk.Label(row,text='Tiger',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left',padx=(15,5))
        tvar=tk.StringVar(value='A')
        ttk.Combobox(row,textvariable=tvar,values=CARDS,state='readonly',width=5).pack(side='left')
        def run_search():
            self.pair.set(dvar.get()+tvar.get())
            self.analyze()
            win.destroy()
        tk.Button(win,text='SEARCH / ANALYZE',command=run_search,bg='#0758d9',fg='white',
                  font=('Segoe UI',10,'bold'),relief='groove',bd=1,padx=16,pady=7).pack(pady=18)
    def open_production_entry(self):
        win=tk.Toplevel(self.root)
        win.title('Production Result Entry')
        win.geometry('720x240')
        win.configure(bg=self._theme()['root'])
        t=self._theme()
        tk.Label(win,text='PRODUCTION RESULT ENTRY',bg=t['root'],fg=t['accent'],font=('Segoe UI',14,'bold')).pack(pady=10)
        row=tk.Frame(win,bg=t['root']); row.pack(fill='x',padx=20,pady=8)
        tk.Label(row,text='Prediction',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left')
        tk.Label(row,textvariable=self.current_prediction,bg=t['root'],fg=t['accent'],font=('Segoe UI',11,'bold')).pack(side='left',padx=8)
        tk.Label(row,text='Came?',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left',padx=(20,4))
        ttk.Combobox(row,textvariable=self.prediction_correct,values=['YES','NO'],state='readonly',width=8).pack(side='left')
        tk.Label(row,text='D:',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left',padx=(15,3))
        ttk.Combobox(row,textvariable=self.actual_dragon,values=CARDS,state='readonly',width=5).pack(side='left')
        tk.Label(row,text='T:',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left',padx=(10,3))
        ttk.Combobox(row,textvariable=self.actual_tiger,values=CARDS,state='readonly',width=5).pack(side='left')
        tk.Label(row,text='Final:',bg=t['root'],fg=t['text'],font=('Segoe UI',10,'bold')).pack(side='left',padx=(10,3))
        ttk.Combobox(row,textvariable=self.final_result,values=RESULTS,state='readonly',width=7).pack(side='left')
        tk.Button(win,text='SAVE RESULT',command=lambda:(self.save_production_result(),win.destroy()),
                  bg='#15803d',fg='white',font=('Segoe UI',10,'bold'),relief='flat',padx=15,pady=7).pack(pady=12)
        tk.Button(win,text='DOWNLOAD HISTORY',command=self.download_history,
                  bg=t['button'],fg=t['text'],font=('Segoe UI',9,'bold'),relief='flat',padx=12,pady=6).pack()

    def refresh_prediction_dashboard(self):
        if not hasattr(self,'prediction_dashboard'):
            return
        for w in self.prediction_dashboard.winfo_children():
            w.destroy()
        bg='#071a35'; green='#00e676'; red='#ff2028'; white='#f8fafc'; yellow='#ffe600'
        total=len(self.result_history)
        matches=sum(1 for x in self.result_history if str(x.get('came','')).upper()=='YES')
        nomatch=max(0,total-matches)
        mp=(matches/total*100) if total else 0
        np=(nomatch/total*100) if total else 0

        left=tk.Frame(self.prediction_dashboard,bg=bg); left.pack(side='left',fill='both',expand=True)
        tk.Label(left,text='✓',bg=green,fg='white',font=('Segoe UI',20,'bold'),width=2).pack(side='left',padx=(22,10),pady=10)
        mbox=tk.Frame(left,bg=bg); mbox.pack(side='left',fill='y',padx=(0,25))
        tk.Label(mbox,text='% MATCH',bg=bg,fg=white,font=('Segoe UI',11,'bold')).pack(anchor='w')
        tk.Label(mbox,text=f'{mp:.0f}%',bg=bg,fg=green,font=('Segoe UI',26,'bold')).pack(anchor='w')
        tk.Label(mbox,text=f'({matches} / {total})',bg=bg,fg=white,font=('Segoe UI',10,'bold')).pack(anchor='w')

        sep=tk.Frame(self.prediction_dashboard,bg='#1e4f83',width=1); sep.pack(side='left',fill='y',pady=8)

        mid=tk.Frame(self.prediction_dashboard,bg=bg); mid.pack(side='left',fill='both',expand=True)
        tk.Label(mid,text='✕',bg=red,fg='white',font=('Segoe UI',20,'bold'),width=2).pack(side='left',padx=(25,10),pady=10)
        nbox=tk.Frame(mid,bg=bg); nbox.pack(side='left',fill='y',padx=(0,25))
        tk.Label(nbox,text='% NO MATCH',bg=bg,fg=white,font=('Segoe UI',11,'bold')).pack(anchor='w')
        tk.Label(nbox,text=f'{np:.0f}%',bg=bg,fg=red,font=('Segoe UI',26,'bold')).pack(anchor='w')
        tk.Label(nbox,text=f'({nomatch} / {total})',bg=bg,fg=white,font=('Segoe UI',10,'bold')).pack(anchor='w')

        donut=tk.Canvas(self.prediction_dashboard,width=170,height=120,bg=bg,highlightthickness=0)
        donut.pack(side='left',padx=10)
        cx,cy=60,60; r=46
        if total:
            donut.create_arc(cx-r,cy-r,cx+r,cy+r,start=90,extent=-360*matches/total,fill=green,outline=bg,width=1)
            donut.create_arc(cx-r,cy-r,cx+r,cy+r,start=90-360*matches/total,extent=-360*nomatch/total,fill=red,outline=bg,width=1)
        else:
            donut.create_oval(cx-r,cy-r,cx+r,cy+r,outline='#334155',width=18)
        donut.create_text(cx,cy-6,text=str(total),fill=white,font=('Segoe UI',17,'bold'))
        donut.create_text(cx,cy+15,text='Total\\nRounds',fill=white,font=('Segoe UI',8,'bold'))

        legend=tk.Frame(self.prediction_dashboard,bg=bg); legend.pack(side='left',fill='y',padx=4)
        tk.Label(legend,text=f'●  Match ({matches})     {mp:.0f}%',bg=bg,fg=green,font=('Segoe UI',10,'bold')).pack(anchor='w',pady=5)
        tk.Label(legend,text=f'●  No Match ({nomatch})     {np:.0f}%',bg=bg,fg=red,font=('Segoe UI',10,'bold')).pack(anchor='w',pady=5)

        recent=tk.Frame(self.prediction_dashboard,bg=bg,highlightthickness=1,highlightbackground='#0b72ff')
        recent.pack(side='right',fill='both',expand=True,padx=(18,8))
        tk.Label(recent,text='RECENT PREDICTION HISTORY (LATEST 10)',bg=bg,fg=white,font=('Segoe UI',10,'bold')).pack(anchor='w',padx=12,pady=(7,2))
        circles=tk.Frame(recent,bg=bg); circles.pack(fill='x',pady=3)
        items=list(reversed(self.result_history[-10:]))
        if not items:
            tk.Label(circles,text='No prediction history yet',bg=bg,fg='#94a3b8',font=('Segoe UI',10,'bold')).pack(pady=16)
        for item in items:
            pred=normalize_result(item.get('prediction','')) or '-'
            came=str(item.get('came','')).upper()=='YES'
            col=green if pred=='D' else ('#ffd500' if pred=='T' else '#8b5cf6')
            cv=tk.Canvas(circles,width=40,height=45,bg=bg,highlightthickness=0)
            cv.create_oval(3,2,37,36,fill=col,outline='')
            cv.create_text(20,19,text=pred,fill='white',font=('Segoe UI',9,'bold'))
            cv.pack(side='left',padx=3)
        status=tk.Frame(recent,bg=bg); status.pack(fill='x')
        for item in items:
            came=str(item.get('came','')).upper()=='YES'
            cv=tk.Canvas(status,width=40,height=25,bg=bg,highlightthickness=0)
            cv.create_oval(10,2,30,22,fill=green if came else red,outline='')
            cv.create_text(20,12,text='✓' if came else '✕',fill='white',font=('Segoe UI',8,'bold'))
            cv.pack(side='left',padx=3)

    def _render_production_table(self):
        if not hasattr(self,'production_tree'):
            return
        for x in self.production_tree.get_children():
            self.production_tree.delete(x)
        items=list(reversed(self.result_history[-8:]))
        for i,item in enumerate(items,1):
            final=normalize_result(item.get('final',''))
            came=str(item.get('came','')).upper()=='YES'
            pred=normalize_result(item.get('prediction',''))
            d=clean_card(item.get('dragon','')); t=clean_card(item.get('tiger',''))
            tie='-' if final!='TIE' else 'TIE'
            pair=item.get('pair','') or (d+t)
            dq=self.oe(d) if d else '-'; tq=self.oe(t) if t else '-'
            prev=item.get('prev_result','-')
            vals=(i,item.get('round_id',''),item.get('round_time','') or item.get('time',''),d,t,tie,pred,final,pair,dq,tq,prev)
            self.production_tree.insert('', 'end', values=vals, tags=('match' if came else 'nomatch',))

    def open_live_browser(self):
        """Ask for a live-game URL, then open it in the Edge/WebView window."""
        url = simpledialog.askstring(
            'LIVE DRAGON TIGER',
            'Enter live game URL:',
            initialvalue=self.live_url,
            parent=self.root
        )
        if not url:
            return
        url = url.strip()
        if not url.lower().startswith(('http://', 'https://')):
            messagebox.showerror('Invalid URL', 'Please enter a valid http:// or https:// URL.')
            return
        self.live_url = url
        if self.live_open:
            self.status.set('LIVE DRAGON TIGER browser is already open.')
            return

        try:
            import webview
        except Exception:
            webbrowser.open(url)
            self.status.set('LIVE browser opened in the default Windows browser.')
            return

        try:
            self.root.update_idletasks()
            x = max(0, self.root.winfo_x() + self.root.winfo_width() + 8)
            y = max(0, self.root.winfo_y())
            w = max(560, min(760, self.root.winfo_width()))
            h = max(650, self.root.winfo_height())
        except Exception:
            x, y, w, h = 100, 80, 700, 720

        self.live_open = True
        self.status.set('Opening LIVE DRAGON TIGER browser beside the analyzer...')

        def run_webview():
            try:
                self.live_window = webview.create_window(
                    'LIVE DRAGON TIGER',
                    url=url,
                    x=x, y=y, width=w, height=h,
                    resizable=True, min_size=(520, 600)
                )
                webview.start(gui='edgechromium', debug=False)
            except Exception:
                self.live_open = False
                self.live_window = None
                try:
                    self.root.after(0, lambda: self.status.set('Embedded live browser unavailable; opened default browser.'))
                    webbrowser.open(url)
                except Exception:
                    pass
            finally:
                self.live_open = False
                self.live_window = None

        threading.Thread(target=run_webview, daemon=True).start()

    def oe(self,c): return 'ODD' if c in ODD else 'EVEN'
    def tag_for(self,r): return {'D':'dragon','T':'tiger','TIE':'tie'}[r]

    def add_result(self):
        n=max([int(r['sno']) for r in self.data if str(r.get('sno','')).isdigit()] or [0])+1
        self.data.append({'sno':n,'round_id':'','time':'','dragon':self.dr.get(),'tiger':self.ti.get(),'result':self.re.get(),'date':''})
        self.refresh()

    def clear_all(self):
        if messagebox.askyesno('Confirm','Delete all loaded rounds?'):
            self.data=[]; self.refresh()

    def refresh(self, rows=None):
        rows = self.data if rows is None else rows
        for x in self.tree.get_children(): self.tree.delete(x)
        # Always display the latest/end result on the LEFT (newest first).
        for idx in range(len(rows)-1,-1,-1):
            r=rows[idx]
            prev = rows[idx-1].get('result','') if idx > 0 else ''
            self.tree.insert('', 'end', values=(r.get('sno',''),r.get('round_id',''),r.get('time',''),r.get('dragon',''),r.get('tiger',''),r.get('result',''),r.get('dragon','')+r.get('tiger',''),self.oe(r.get('dragon','')),self.oe(r.get('tiger','')),prev),tags=(self.tag_for(r.get('result','D')),))
        for x in self.last_frame.winfo_children(): x.destroy()
        recent = list(reversed(self.data[-14:]))
        for r in recent:
            res=r.get('result','')
            card=f'{r.get("dragon","")}/{r.get("tiger","")}'
            color=RESULT_COLORS.get(res,'#374151')
            box=tk.Canvas(self.last_frame,width=58,height=78,bg='#111827',highlightthickness=0)
            box.create_oval(5,3,53,51,fill=color,outline='')
            box.create_text(29,27,text=res,fill='white',font=('Segoe UI',11,'bold'))
            box.create_text(29,64,text=card,fill='white',font=('Segoe UI',8,'bold'))
            box.pack(side='right',padx=2)
        self.update_last_analysis()
        self._render_production_table()
        self.refresh_prediction_dashboard()
        self.update_title_counts()

    def update_last_analysis(self):
        if not self.data:
            self.last_analysis.set('LAST RESULT ANALYSIS | No rounds loaded')
            return
        idx=len(self.data)-1
        cur=self.data[idx]
        cur_pair=f'{cur.get("dragon","")}{cur.get("tiger","")}'
        cur_desc=f'{cur.get("result","")} | D:{cur.get("dragon","")} ({self.oe(cur.get("dragon",""))}) | T:{cur.get("tiger","")} ({self.oe(cur.get("tiger",""))}) | Pair:{cur_pair}'
        prev=[]
        for j in range(max(0,idx-5),idx):
            q=self.data[j]
            prev.append(f'{q.get("result","")}[{q.get("dragon","")}/{q.get("tiger","")}]')
        # Historical occurrences of the current result + pair: show what came immediately after each occurrence.
        matches=[]
        for j,r in enumerate(self.data[:-1]):
            if r.get("result")==cur.get("result") and r.get("dragon","")+r.get("tiger","")==cur_pair:
                nxt=self.data[j+1]
                matches.append(f'{nxt.get("result","")} D:{nxt.get("dragon","")} T:{nxt.get("tiger","")}')
        after=' | '.join(matches[-3:]) if matches else 'No previous matching occurrence with a following round'
        self.last_analysis.set(
            f'LAST: {cur_desc}   |   PREVIOUS 5: {"  ".join(prev) if prev else "-"}   |   AFTER SAME RESULT+PAIR: {after}'
        )

    def update_title_counts(self):
        c=Counter(r.get('result') for r in self.data)
        self.summary.set(f'PAIR {self.pair.get().upper()} | D {c["D"]} | T {c["T"]} | TIE {c["TIE"]} | TOTAL {len(self.data)}')

    def analyze(self):
        p=self.pair.get().upper().strip()
        # Record every Pair Analysis search without changing the existing statistical analysis.
        search_time=datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        existing_rows=[r for r in self.data if r.get('dragon','')+r.get('tiger','')==p]
        search_sequence=[r.get('result','') for r in existing_rows if r.get('result','') in RESULTS]
        self.pair_history.append({
            'time':search_time,'pair':p,
            'prediction': '',
            'sequence': search_sequence[-30:]
        })
        self.pair_history=self.pair_history[-100:]
        self.render_pair_history()
        self.save_settings()

        rows=existing_rows
        cnt=Counter(r.get('result') for r in rows)
        self.report_rows=[]
        if not rows:
            self.render_report_rows()
            self.summary.set(f'PAIR {p} | NO HISTORY')
            self.current_prediction.set('NO PREDICTION')
            self.prediction_pct.set('')
            if self.pair_history:
                self.pair_history[-1]['prediction']=''
                self.render_pair_history(); self.save_settings()
            return
        maxres=max(RESULTS,key=lambda x:cnt[x]); pct=cnt[maxres]/len(rows)*100
        self.summary.set(f'PAIR {p} | CAME {len(rows)} TIMES | D {cnt["D"]} | T {cnt["T"]} | TIE {cnt["TIE"]} | MOST {maxres} ({pct:.1f}%)')
        self.current_prediction.set(maxres)
        self.prediction_pct.set(f'{pct:.1f}%')
        if self.pair_history:
            self.pair_history[-1]['prediction']=maxres
            self.render_pair_history(); self.save_settings()
        self.draw_frequency_chart(rows)

        for idx,r in enumerate(self.data):
            if r.get('dragon','')+r.get('tiger','') != p: continue
            prev=[]
            for j in range(max(0,idx-3),idx):
                q=self.data[j]; prev.append(f'{q.get("result","")} {q.get("dragon","")}{q.get("tiger","")}')
            nxt=[]
            for j in range(idx+1,min(len(self.data),idx+7)):
                q=self.data[j]; nxt.append(f'{q.get("result","")} {q.get("dragon","")}{q.get("tiger","")}')
            while len(nxt)<6: nxt.append('-')
            self.report_rows.append((r.get('sno',''),' | '.join(prev) if prev else '-',f'{r.get("result","")} {p}',nxt[0],nxt[1],nxt[2],nxt[3],nxt[4],nxt[5]))
        self.render_report_rows()

    def _report_circle(self, cv, x, y, result, diameter=20):
        result=normalize_result(result)
        if not result:
            return x
        color=RESULT_COLORS.get(result,'#64748b')
        cv.create_oval(x,y,x+diameter,y+diameter,fill=color,outline='')
        label='TIE' if result=='TIE' else result
        fsize=6 if result=='TIE' else 8
        cv.create_text(x+diameter/2,y+diameter/2,text=label,fill='white',
                       font=('Segoe UI',fsize,'bold'))
        return x+diameter

    def _draw_report_cell(self, cv, x, y, width, height, value, align='center'):
        value=str(value or '').strip()
        if not value:
            return
        if value=='-':
            cv.create_text(x+width/2,y+height/2,text='-',fill='white',font=('Segoe UI',8,'bold'))
            return
        parts=[p.strip() for p in value.split('|')]
        cursor=x+6
        for n,part in enumerate(parts):
            m=re.match(r'^(D|T|TIE)\\s*(.*)$',part,re.I)
            if m:
                result=normalize_result(m.group(1))
                rest=m.group(2).strip()
                cursor=self._report_circle(cv,cursor,y+(height-20)/2,result,20)+5
                if rest:
                    cv.create_text(cursor,y+height/2,text=rest,anchor='w',
                                   fill='white',font=('Segoe UI',8,'bold'))
                    cursor += max(22,len(rest)*5.3)
            else:
                cv.create_text(cursor,y+height/2,text=part,anchor='w',
                               fill='white',font=('Segoe UI',8,'bold'))
                cursor += max(20,len(part)*5.3)
            if n < len(parts)-1:
                cv.create_text(cursor,y+height/2,text='|',anchor='w',
                               fill='#cbd5e1',font=('Segoe UI',8,'bold'))
                cursor += 10

    def render_report_rows(self):
        if not hasattr(self,'report_canvas'):
            return
        cv=self.report_canvas
        cv.delete('all')
        if not getattr(self,'report_rows',None):
            cv.configure(scrollregion=(0,0,max(1,cv.winfo_width()),1))
            return
        base=[90,250,90,90,90,90,90,90,90]
        total=sum(base)
        W=max(total,cv.winfo_width())
        scale=W/total
        widths=[w*scale for w in base]
        row_h=38
        for ridx,row in enumerate(self.report_rows):
            y=ridx*row_h
            res=normalize_result(str(row[2]).split()[0])
            row_bg={'D':'#14532d','T':'#92400e','TIE':'#4c1d95'}.get(res,'#12304f')
            cv.create_rectangle(0,y,W,y+row_h,fill=row_bg,outline='#0b4f72')
            x=0
            for cidx,(val,w) in enumerate(zip(row,widths)):
                if cidx==0:
                    cv.create_text(x+w/2,y+row_h/2,text=str(val),fill='white',
                                   font=('Segoe UI',8,'bold'))
                else:
                    self._draw_report_cell(cv,x,y,w,row_h,val)
                x+=w
        cv.configure(scrollregion=(0,0,W,len(self.report_rows)*row_h))

    def _circle(self, parent, result, size=40):
        result=normalize_result(result) or 'TIE'
        color=RESULT_COLORS.get(result,'#7c3aed')
        cv=tk.Canvas(parent,width=size,height=size,bg=parent.cget('bg'),highlightthickness=0,cursor='hand2')
        pad=2
        cv.create_oval(pad,pad,size-pad,size-pad,fill=color,outline='')
        cv.create_text(size/2,size/2,text=result,fill='white',
                       font=('Segoe UI',max(8,int(size*0.25)),'bold'))
        return cv

    def _show_pair_result_popup(self, item):
        pair=str(item.get('pair','')).strip().upper().replace(' ','')
        pred=normalize_result(item.get('prediction','')) or 'TIE'
        d=clean_card(pair[:2] if len(pair)==3 and pair[:2]=='10' else pair[:1])
        t=clean_card(pair[2:] if len(pair)==3 and pair[:2]=='10' else pair[1:2])
        if not d or not t:
            d=clean_card(item.get('dragon',''))
            t=clean_card(item.get('tiger',''))

        # Compact Pair Result popup with enough height for DELETE RESULT and CLOSE.
        win=tk.Toplevel(self.root)
        win.title('Pair Result')
        win.geometry('220x310')
        win.resizable(False,False)
        win.configure(bg='white')
        try:
            x=self.root.winfo_pointerx()+12; y=self.root.winfo_pointery()+12
            win.geometry(f'220x310+{x}+{y}')
        except Exception:
            pass

        tk.Label(win,text='Dragon',bg='white',fg='#111111',
                 font=('Segoe UI',13,'normal')).pack(pady=(6,1))
        dcard=tk.Frame(win,bg='white',highlightthickness=1,highlightbackground='#ffe600')
        dcard.pack(pady=(0,1))
        tk.Label(dcard,text=d or '-',bg='white',fg='#e11d2e',
                 font=('Segoe UI',19,'bold'),width=2,height=1).pack(padx=4,pady=1)
        tk.Label(win,text='Tiger',bg='white',fg='#111111',
                 font=('Segoe UI',13,'normal')).pack(pady=(0,1))

        lower=tk.Frame(win,bg='white')
        lower.pack(pady=(0,2))
        tcard=tk.Frame(lower,bg='white',highlightthickness=1,highlightbackground='#ffe600')
        tcard.pack(side='left',padx=(0,6))
        tk.Label(tcard,text=t or '-',bg='white',fg='#111111',
                 font=('Segoe UI',19,'bold'),width=2,height=1).pack(padx=4,pady=1)

        trophy=tk.Canvas(lower,width=53,height=53,bg='white',highlightthickness=0)
        trophy.pack(side='left')
        trophy.create_oval(2,2,51,51,fill='#eef2e8',outline='')
        trophy.create_rectangle(19,15,34,34,fill='#16a34a',outline='')
        trophy.create_arc(12,13,23,30,start=90,extent=180,style='arc',outline='#16a34a',width=5)
        trophy.create_arc(30,13,41,30,start=-90,extent=180,style='arc',outline='#16a34a',width=5)
        trophy.create_rectangle(16,33,37,37,fill='#16a34a',outline='')
        trophy.create_rectangle(23,37,30,42,fill='#16a34a',outline='')
        trophy.create_rectangle(17,42,36,45,fill='#16a34a',outline='')
        trophy.create_text(27,27,text='★',fill='white',font=('Segoe UI',7,'bold'))

        result_color=RESULT_COLORS.get(pred,'#7c3aed')
        tk.Label(win,text=f'Result: {pred}',bg='white',fg=result_color,
                 font=('Segoe UI',8,'bold')).pack(pady=(1,1))
        tk.Label(win,text=f'Dragon {self.oe(d) if d else "-"} | Tiger {self.oe(t) if t else "-"}',
                 bg='white',fg='#475569',font=('Segoe UI',6,'bold')).pack(pady=1)
        def delete_this_result():
            try:
                idx=next((i for i,x in enumerate(self.pair_history) if x is item), -1)
                if idx < 0:
                    idx=next((i for i,x in enumerate(self.pair_history)
                              if x.get('time')==item.get('time') and x.get('pair')==item.get('pair')
                              and x.get('prediction')==item.get('prediction')), -1)
                if idx >= 0:
                    del self.pair_history[idx]
                    self.save_settings()
                    self.render_pair_history()
                win.destroy()
            except Exception as e:
                messagebox.showerror('Delete Result',str(e))

        tk.Button(win,text='DELETE RESULT',command=delete_this_result,bg='#b91c1c',fg='white',
                  activebackground='#dc2626',activeforeground='white',
                  font=('Segoe UI',9,'bold'),relief='groove',bd=1,padx=12,pady=4).pack(pady=(4,3))
        tk.Button(win,text='CLOSE',command=win.destroy,bg='#0758d9',fg='white',
                  activebackground='#0b78ff',activeforeground='white',
                  font=('Segoe UI',9,'bold'),relief='groove',bd=1,padx=12,pady=4).pack(pady=2)
        win.transient(self.root)
        win.lift()

    def render_pair_history(self):
        if not hasattr(self,'pair_history_frame'): return
        for w in self.pair_history_frame.winfo_children(): w.destroy()
        items=list(reversed(self.pair_history[-15:]))
        if not items:
            tk.Label(self.pair_history_frame,text='PAIR SEARCH HISTORY WILL APPEAR HERE',bg='#111827',fg='#9ca3af',font=('Segoe UI',9,'bold')).pack(pady=25)
            return

        # Show exactly the latest 15 searches as one continuous compact row.
        # Smaller circles keep the full 15 visible; click any circle for pair/card odd-even details.
        strip=tk.Frame(self.pair_history_frame,bg='#111827')
        strip.pack(side='left',fill='x',expand=True)
        for item in items:
            slot=tk.Frame(strip,bg='#111827',width=52,height=50,cursor='hand2')
            slot.pack(side='left',fill='y',padx=0)
            slot.pack_propagate(False)
            pred=normalize_result(item.get('prediction','')) or 'TIE'
            cv=self._circle(slot,pred,40)
            cv.pack(anchor='center',pady=2)
            cv.bind('<Button-1>',lambda e,it=item:self._show_pair_result_popup(it))
            slot.bind('<Button-1>',lambda e,it=item:self._show_pair_result_popup(it))

        tk.Label(self.pair_history_frame,text=f'Total History: {len(self.pair_history)}',
                 bg='#111827',fg='#ffe600',font=('Segoe UI',8,'bold'),
                 relief='groove',bd=1,padx=8,pady=4).pack(side='right',padx=(4,2),pady=8)

    def save_production_result(self):
        correct=self.prediction_correct.get().strip().upper()
        d=clean_card(self.actual_dragon.get())
        t=clean_card(self.actual_tiger.get())
        final=normalize_result(self.final_result.get())
        prediction=normalize_result(self.current_prediction.get())
        if correct not in ('YES','NO'):
            messagebox.showwarning('Save Result','Select YES or NO for whether the prediction came.')
            return
        if not d or not t:
            messagebox.showwarning('Save Result','Select both the actual D and T card numbers.')
            return
        if not final:
            messagebox.showwarning('Save Result','Select the final result.')
            return
        latest=self.data[-1] if self.data else {}
        record={
            'time':datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
            'round_id':str(latest.get('round_id','')),
            'round_time':str(latest.get('time','')),
            'pair':self.pair.get().upper().strip() or (d+t),
            'prediction':prediction,
            'came':correct,
            'dragon':d,'tiger':t,'final':final,
            'prev_result':str(self.data[-2].get('result','')) if len(self.data)>1 else '-'
        }
        self.result_history.append(record)
        self.result_history=self.result_history[-200:]
        self.save_settings()
        self.prediction_correct.set(''); self.actual_dragon.set(''); self.actual_tiger.set('')
        messagebox.showinfo('Saved','Production result saved successfully.')
        self.show_result_history()


    def show_pair_history(self):
        win=tk.Toplevel(self.root); win.title('Pair Search History'); win.geometry('900x500'); win.configure(bg=self._theme()['root'])
        t=self._theme()
        tk.Label(win,text='PAIR SEARCH HISTORY',bg=t['root'],fg=t['accent'],font=('Segoe UI',14,'bold')).pack(pady=10)
        frame=tk.Frame(win,bg=t['root']); frame.pack(fill='both',expand=True,padx=12,pady=5)
        cols=('NO','TIME','PAIR','PREDICTION','RESULT SEQUENCE','DELETE')
        tree=ttk.Treeview(frame,columns=cols,show='headings',height=15)
        widths={'NO':50,'TIME':150,'PAIR':80,'PREDICTION':100,'RESULT SEQUENCE':420,'DELETE':80}
        for col in cols: tree.heading(col,text=col); tree.column(col,width=widths[col],anchor='center')
        tree.pack(fill='both',expand=True)
        for i,item in enumerate(reversed(self.pair_history),1):
            seq=' '.join(item.get('sequence',[]))
            tree.insert('', 'end',iid=str(i-1),values=(i,item.get('time',''),item.get('pair',''),item.get('prediction',''),seq,'DELETE'))
        def delete_selected():
            sel=tree.selection()
            if not sel: return
            rev_index=int(sel[0]); actual_index=len(self.pair_history)-1-rev_index
            if 0<=actual_index<len(self.pair_history):
                del self.pair_history[actual_index]; self.save_settings(); self.render_pair_history()
                tree.delete(sel[0])
                # Rebuild IDs/order after deletion.
                for x in tree.get_children(): tree.delete(x)
                for i,item in enumerate(reversed(self.pair_history),1):
                    tree.insert('', 'end',iid=str(i-1),values=(i,item.get('time',''),item.get('pair',''),item.get('prediction',''),' '.join(item.get('sequence',[])),'DELETE'))
        tk.Button(win,text='DELETE SELECTED',command=delete_selected,bg='#b91c1c',fg='white',font=('Segoe UI',9,'bold'),relief='flat',padx=12,pady=6).pack(side='left',padx=12,pady=8)
        tk.Button(win,text='DOWNLOAD HISTORY',command=self.download_history,bg='#374151',fg='white',font=('Segoe UI',9,'bold'),relief='flat',padx=12,pady=6).pack(side='right',padx=12,pady=8)

    def show_result_history(self):
        win=tk.Toplevel(self.root); win.title('Production Result History'); win.geometry('1050x520'); win.configure(bg=self._theme()['root'])
        t=self._theme()
        tk.Label(win,text='PRODUCTION RESULT HISTORY — LATEST 15+',bg=t['root'],fg=t['accent'],font=('Segoe UI',14,'bold')).pack(pady=10)
        frame=tk.Frame(win,bg=t['root']); frame.pack(fill='both',expand=True,padx=12,pady=5)
        cols=('NO','TIME','PAIR','PREDICTION','CAME?','D','T','FINAL','DELETE')
        tree=ttk.Treeview(frame,columns=cols,show='headings',height=15)
        widths={'NO':50,'TIME':150,'PAIR':80,'PREDICTION':100,'CAME?':80,'D':60,'T':60,'FINAL':80,'DELETE':80}
        for col in cols: tree.heading(col,text=col); tree.column(col,width=widths[col],anchor='center')
        for i,item in enumerate(reversed(self.result_history),1):
            tree.insert('', 'end',iid=str(i-1),values=(i,item.get('time',''),item.get('pair',''),item.get('prediction',''),item.get('came',''),item.get('dragon',''),item.get('tiger',''),item.get('final',''),'DELETE'))
        tree.pack(fill='both',expand=True)
        def delete_selected():
            sel=tree.selection()
            if not sel: return
            rev_index=int(sel[0]); actual_index=len(self.result_history)-1-rev_index
            if 0<=actual_index<len(self.result_history):
                del self.result_history[actual_index]; self.save_settings()
                for x in tree.get_children(): tree.delete(x)
                for i,item in enumerate(reversed(self.result_history),1):
                    tree.insert('', 'end',iid=str(i-1),values=(i,item.get('time',''),item.get('pair',''),item.get('prediction',''),item.get('came',''),item.get('dragon',''),item.get('tiger',''),item.get('final',''),'DELETE'))
        tk.Button(win,text='DELETE SELECTED',command=delete_selected,bg='#b91c1c',fg='white',font=('Segoe UI',9,'bold'),relief='flat',padx=12,pady=6).pack(side='left',padx=12,pady=8)
        tk.Button(win,text='DOWNLOAD HISTORY',command=self.download_history,bg='#374151',fg='white',font=('Segoe UI',9,'bold'),relief='flat',padx=12,pady=6).pack(side='right',padx=12,pady=8)

    def download_history(self):
        if not self.pair_history and not self.result_history:
            messagebox.showwarning('Download History','No saved history is available.')
            return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',initialfile='DragonTiger_History.xlsx',filetypes=[('Excel','*.xlsx'),('CSV','*.csv')])
        if not path: return
        try:
            if path.lower().endswith('.csv'):
                rows=[['TYPE','TIME','PAIR','PREDICTION','CAME?','D','T','FINAL','RESULT SEQUENCE']]
                for x in reversed(self.pair_history):
                    rows.append(['PAIR SEARCH',x.get('time',''),x.get('pair',''),x.get('prediction',''),'','','','', ' '.join(x.get('sequence',[]))])
                for x in reversed(self.result_history):
                    rows.append(['PRODUCTION RESULT',x.get('time',''),x.get('pair',''),x.get('prediction',''),x.get('came',''),x.get('dragon',''),x.get('tiger',''),x.get('final',''),''])
                with open(path,'w',newline='',encoding='utf-8-sig') as f: csv.writer(f).writerows(rows)
            else:
                from openpyxl import Workbook
                wb=Workbook()
                ws=wb.active; ws.title='PAIR SEARCH HISTORY'
                ws.append(['TIME','PAIR','PREDICTION','RESULT SEQUENCE'])
                for x in reversed(self.pair_history): ws.append([x.get('time',''),x.get('pair',''),x.get('prediction',''),' '.join(x.get('sequence',[]))])
                ws2=wb.create_sheet('PRODUCTION RESULTS')
                ws2.append(['TIME','PAIR','PREDICTION','CAME?','D','T','FINAL RESULT'])
                for x in reversed(self.result_history): ws2.append([x.get('time',''),x.get('pair',''),x.get('prediction',''),x.get('came',''),x.get('dragon',''),x.get('tiger',''),x.get('final','')])
                wb.save(path)
            messagebox.showinfo('Download History','History downloaded successfully.')
        except Exception as e:
            messagebox.showerror('Download History',str(e))

    def draw_frequency_chart(self, rows):
        for w in self.chart_frame.winfo_children(): w.destroy()
        if not rows:
            tk.Label(self.chart_frame,text='No matching pair history',bg='#111827',fg='#9ca3af').pack(pady=25)
            return
        counts=Counter()
        for r in rows:
            d,t=r.get('dragon',''),r.get('tiger','')
            if d: counts[(d,'D')]+=1
            if t: counts[(t,'T')]+=1
        top=sorted(counts.items(),key=lambda x:(-x[1],x[0][0],x[0][1]))[:13]
        maxv=max(v for _,v in top) or 1
        for (num,side),v in top:
            row=tk.Frame(self.chart_frame,bg='#111827'); row.pack(fill='x',pady=2)
            tk.Label(row,text=f'{num}  {side}',bg='#14532d' if side=='D' else '#92400e',fg='white',width=7,font=('Segoe UI',9,'bold')).pack(side='left')
            bar=tk.Canvas(row,height=20,bg='#1f2937',highlightthickness=0); bar.pack(side='left',fill='x',expand=True,padx=6)
            bar.update_idletasks()
            w=max(30,int(bar.winfo_width()*v/maxv))
            bar.create_rectangle(0,2,w,18,fill='#22c55e' if side=='D' else '#f59e0b',outline='')
            tk.Label(row,text=f'{v} times',bg='#111827',fg='white',width=10,anchor='e',font=('Segoe UI',9,'bold')).pack(side='right')

    def draw_round_graph(self):
        cv=self.round_chart
        cv.delete('all')
        data=self.data[-40:]
        if not data:
            cv.create_text(20,80,anchor='w',text='Import data to view round analysis',fill='#9ca3af',font=('Segoe UI',10))
            return
        cv.update_idletasks()
        W=max(cv.winfo_width(),700); H=180
        left=45; right=20; top=18; bottom=38
        usable=max(1,W-left-right)
        step=usable/max(1,len(data))
        barw=max(6,min(18,step*0.65))
        cv.create_line(left,H-bottom,left+usable,H-bottom,fill='#4b5563')
        for i,r in enumerate(data):
            x=left+i*step+step/2
            res=r.get('result','')
            h={'D':85,'T':65,'TIE':45}.get(res,35)
            color={'D':'#22c55e','T':'#f59e0b','TIE':'#a855f7'}.get(res,'#6b7280')
            cv.create_rectangle(x-barw/2,H-bottom-h,x+barw/2,H-bottom,fill=color,outline='')
            if len(data)<=20 or i%2==0:
                cv.create_text(x,H-bottom-h-9,text=res,fill='white',font=('Segoe UI',8,'bold'))
            label=str(r.get('sno',''))
            if len(data)<=25 or i%4==0:
                cv.create_text(x,H-bottom+12,text=label,fill='#9ca3af',font=('Segoe UI',8))
            num=f"{r.get('dragon','')}/{r.get('tiger','')}"
            if len(data)<=20 or i%3==0:
                cv.create_text(x,H-bottom+27,text=num,fill='#d1d5db',font=('Segoe UI',7))
        cv.create_text(left,8,anchor='w',text='Latest 40 rounds | bar = result | label = D/T card numbers',fill='#cbd5e1',font=('Segoe UI',9,'bold'))

    def import_data(self):
        path=filedialog.askopenfilename(filetypes=[('Excel','*.xlsx'),('CSV','*.csv'),('All files','*.*')])
        if not path:return
        try:
            rows=[]
            if path.lower().endswith('.csv'):
                with open(path,newline='',encoding='utf-8-sig') as f: rows=list(csv.reader(f))
            else:
                from openpyxl import load_workbook
                wb=load_workbook(path,data_only=True); ws=wb.active; rows=list(ws.iter_rows(values_only=True))
            out=[]; seen=set()
            for row in rows:
                if len(row)<4: continue
                try: sno=int(row[0])
                except: continue
                d=clean_card(row[1]); t=clean_card(row[2]); res=normalize_result(row[3])
                if d and t and res and sno not in seen:
                    out.append({'sno':sno,'round_id':str(row[4]) if len(row)>4 else '','time':str(row[5]) if len(row)>5 else '','dragon':d,'tiger':t,'result':res,'date':str(row[6]) if len(row)>6 else ''}); seen.add(sno)
            out.sort(key=lambda r:r['sno']); self.data=out; self.refresh(); messagebox.showinfo('Import',f'Imported {len(out)} unique rounds.')
        except Exception as e: messagebox.showerror('Import error',str(e))

    def export_data(self):
        if not self.data: messagebox.showwarning('Export','No data to export.'); return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',filetypes=[('Excel','*.xlsx'),('CSV','*.csv')])
        if not path:return
        try:
            rows=[['S NO','ROUND ID','TIME','DRAGON','TIGER','RESULT','PAIR','DRAGON O/E','TIGER O/E','DATE']]
            for r in reversed(self.data): rows.append([r.get('sno',''),r.get('round_id',''),r.get('time',''),r.get('dragon',''),r.get('tiger',''),r.get('result',''),r.get('dragon','')+r.get('tiger',''),self.oe(r.get('dragon','')),self.oe(r.get('tiger','')),r.get('date','')])
            if path.lower().endswith('.csv'):
                with open(path,'w',newline='',encoding='utf-8-sig') as f: csv.writer(f).writerows(rows)
            else:
                from openpyxl import Workbook
                wb=Workbook(); ws=wb.active; ws.title='ANALYZER'
                for row in rows: ws.append(row)
                wb.save(path)
            messagebox.showinfo('Export','Export completed successfully.')
        except Exception as e: messagebox.showerror('Export error',str(e))

    def parse_pasted_rows(self, text):
        lines=[x.strip() for x in text.replace('\r','').split('\n') if x.strip()]
        out=[]
        for line in lines:
            parts=re.split(r'\t|\s{2,}|\|',line)
            joined=' '.join(parts)
            if re.search(r'round\s*id|match\s*time|winner',joined,re.I):
                continue
            rid=''
            m=re.search(r'\b(\d{10,20})\b',line)
            if m: rid=m.group(1)
            res=normalize_result(re.search(r'\b(DRAGON|TIGER|TIE|DRAW)\b',line,re.I).group(1)) if re.search(r'\b(DRAGON|TIGER|TIE|DRAW)\b',line,re.I) else ''
            cards=[]
            for token in re.findall(r'(?<![A-Z0-9])(10|[2-9AJQK])(?![A-Z0-9])',line.upper()): cards.append(token)
            if rid and res:
                d=cards[0] if len(cards)>=2 else ''
                t=cards[1] if len(cards)>=2 else ''
                out.append({'round_id':rid,'time':'','dragon':d,'tiger':t,'result':res,'date':self.collect_date.get()})
        return out

    def merge_collected(self, incoming):
        if not incoming: return 0
        existing_ids={str(r.get('round_id','')) for r in self.data if r.get('round_id')}
        next_sno=max([int(r['sno']) for r in self.data if str(r.get('sno','')).isdigit()] or [0])+1
        added=0
        limit=max(1,int(self.entries.get()))
        for r in incoming[:limit]:
            rid=str(r.get('round_id','')).strip()
            if not rid or rid in existing_ids: continue
            r['sno']=next_sno; next_sno+=1; self.data.append(r); existing_ids.add(rid); added+=1
        self.data.sort(key=lambda x:int(x.get('sno',0)) if str(x.get('sno','')).isdigit() else 0)
        self.refresh()
        return added

    def paste_results(self):
        try:
            text=self.root.clipboard_get()
        except Exception:
            messagebox.showwarning('Paste Results','Copy the YoloBet9 Casino Results rows first, then click PASTE RESULTS.')
            return
        incoming=self.parse_pasted_rows(text)
        if not incoming:
            messagebox.showwarning('Paste Results','No Round ID + Winner rows were detected. Copy the result table rows and try again.')
            return
        added=self.merge_collected(incoming)
        self.status.set(f'Date {self.collect_date.get()} | detected {len(incoming)} | added {added} | duplicates skipped')
        messagebox.showinfo('YoloBet9 Import',f'Detected {len(incoming)} results. Added {added} new unique Round IDs. Duplicate Round IDs were skipped.')

    def filter_date(self):
        target=self.collect_date.get().strip()
        rows=[]
        for r in self.data:
            d=str(r.get('date','')).strip()
            if d==target: rows.append(r)
            elif not d and str(r.get('time','')).startswith(target): rows.append(r)
        self.refresh(rows)
        self.status.set(f'Date filter {target} | showing {len(rows)} stored rounds')


if __name__=='__main__':
    root=tk.Tk(); App(root); root.mainloop()
